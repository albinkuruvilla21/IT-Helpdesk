#!/usr/bin/env python3
"""
Convert everything in inputs/ into text Claude can read, and index it.

  uv run tools/ingest.py            convert new or changed inputs
  uv run tools/ingest.py --all      reconvert everything
  uv run tools/ingest.py --status   report what is new, changed or unchanged without converting

  inputs/brd/          BRD (.xlsx, .docx, .pdf)      → .xlsx: one Markdown table per sheet WITH Excel row
                                                       numbers, so criteria can cite "BRD: Sheet!R12";
                                                       other formats via markitdown
  inputs/transcripts/  PO walkthrough (.vtt, .docx)  → .vtt: "[hh:mm:ss] Speaker: text", one line per turn;
                                                       .docx via markitdown
  inputs/design/       accepted design (.html, .png) → HTML copied with embedded base64 images stubbed out;
                                                       images listed (Claude reads images directly)
  inputs/branding/     brand guidelines (.pdf, …)    → text via markitdown; open the PDF itself for visuals
  inputs/other/        anything else                 → markitdown where it can

Output goes to inputs/_converted/<same subfolder>/<name>.md (or .html), plus inputs/_converted/INDEX.md,
which records each source's SHA-256 so a skill can tell when an input changed after it was converted.
Needs: the packages in pyproject.toml, which uv run installs.
"""
from __future__ import annotations

import argparse
import hashlib
import html
import re
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
INPUTS = ROOT / "inputs"
OUT = INPUTS / "_converted"
INDEX = OUT / "INDEX.md"

MARKITDOWN_EXT = {".xls", ".docx", ".pdf", ".pptx", ".csv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}
COPY_EXT = {".md", ".txt", ".json"}
SKIP_NAMES = {"README.md", ".gitkeep"}


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        for chunk in iter(lambda: f.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()[:16]


def read_index() -> dict[str, tuple[str, str]]:
    """source path → (hash, index row), from the last run."""
    if not INDEX.exists():
        return {}
    known = {}
    for line in INDEX.read_text(encoding="utf-8").splitlines():
        m = re.match(r"^\| `([^`]+)` \| [^|]+ \| `([0-9a-f]{16})` \|", line)
        if m:
            known[m.group(1)] = (m.group(2), line)
    return known


# ─── Converters ──────────────────────────────────────────────────────────────────────────────────
_markitdown = None


def convert_markitdown(src: Path) -> str:
    global _markitdown
    if _markitdown is None:
        try:
            from markitdown import MarkItDown
        except ImportError:
            sys.exit("markitdown is not installed. Run the tools through uv (uv run tools/<tool>.py), or: pip install --group tools")
        _markitdown = MarkItDown()
    return _markitdown.convert(str(src)).text_content


def column_letter(index: int) -> str:
    letters = ""
    index += 1
    while index:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _md_cell(value) -> str:
    if value is None:
        return ""
    return str(value).replace("|", r"\|").replace("\r", "").replace("\n", "<br>").strip()


def convert_xlsx(src: Path) -> str:
    """Every sheet as a Markdown table with Excel's row number first, so a citation can point at a row."""
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl is not installed. Run the tools through uv (uv run tools/<tool>.py), or: pip install --group tools")
    wb = load_workbook(src, data_only=True, read_only=True)
    out = [f"# {src.name}", "", "Column letters and row numbers are Excel's. Cite rows as `BRD: <Sheet>!R<row>`.", ""]
    for ws in wb.worksheets:
        rows = [(i, r) for i, r in enumerate(ws.iter_rows(values_only=True), start=1) if any(v not in (None, "") for v in r)]
        out += [f"## {ws.title}", ""]
        if not rows:
            out += ["(empty)", ""]
            continue
        width = max(len(r) for _, r in rows)
        while width > 1 and all(len(r) < width or r[width - 1] in (None, "") for _, r in rows):
            width -= 1
        out.append("| R | " + " | ".join(column_letter(c) for c in range(width)) + " |")
        out.append("|---" * (width + 1) + "|")
        for i, r in rows:
            values = list(r[:width]) + [None] * (width - len(r[:width]))
            out.append(f"| {i} | " + " | ".join(_md_cell(v) for v in values) + " |")
        out.append("")
    return "\n".join(out)


_TS = re.compile(r"^(\d{1,2}:)?(\d{2}):(\d{2})[.,]\d{3}\s+-->")
_VOICE = re.compile(r"<v(?:\.[^ >]*)?\s+([^>]+)>(.*?)(?:</v>|$)", re.S)


def convert_vtt(src: Path) -> str:
    """WebVTT (Teams, Zoom, Stream) → one line per speaker turn, consecutive cues merged."""
    turns: list[list[str]] = []  # [timestamp, speaker, text]
    stamp = None
    for line in src.read_text(encoding="utf-8-sig", errors="replace").splitlines():
        m = _TS.match(line.strip())
        if m:
            hours = (m.group(1) or "0:").rstrip(":")
            stamp = f"{int(hours):02d}:{m.group(2)}:{m.group(3)}"
            continue
        if stamp is None or not line.strip() or line.strip().isdigit() or line.startswith("NOTE"):
            continue
        v = _VOICE.search(line)
        speaker, text = (v.group(1).strip(), v.group(2)) if v else ("", line)
        if not v and ":" in line[:40] and not line.startswith("http"):
            maybe, rest = line.split(":", 1)
            if 0 < len(maybe) <= 40:
                speaker, text = maybe.strip(), rest
        text = html.unescape(re.sub(r"<[^>]+>", "", text)).strip()
        if not text:
            continue
        if turns and turns[-1][1] == speaker:
            turns[-1][2] += " " + text
        else:
            turns.append([stamp, speaker, text])
    out = [f"# Transcript — {src.name}", "", "Timestamps are the start of each speaker turn. Cite them as `Transcript hh:mm:ss`.", ""]
    for t, speaker, text in turns:
        out.append(f"[{t}] {speaker + ': ' if speaker else ''}{text}")
    return "\n".join(out) + "\n"


_DATA_URI = re.compile(r"data:([a-z]+/[a-z0-9.+-]+);base64,[A-Za-z0-9+/=\s]{200,}", re.I)


def convert_html(src: Path) -> tuple[str, int]:
    """Keep the markup, styles and scripts; replace embedded base64 images with a short stub."""
    text = src.read_text(encoding="utf-8", errors="replace")
    count = 0

    def stub(m):
        nonlocal count
        count += 1
        return f"data:{m.group(1)};base64,EMBEDDED-IMAGE-{count}-REMOVED-SEE-ORIGINAL"

    return _DATA_URI.sub(stub, text), count


# ─── Driver ──────────────────────────────────────────────────────────────────────────────────────
def plan():
    rows = []
    for src in sorted(INPUTS.rglob("*")):
        if not src.is_file() or OUT in src.parents or src.name in SKIP_NAMES or src.name.startswith("~$"):
            continue
        rel = src.relative_to(INPUTS).as_posix()
        ext = src.suffix.lower()
        stem_out = OUT / src.relative_to(INPUTS).parent / src.stem
        if ext == ".vtt":
            kind, dest = "transcript", stem_out.with_suffix(".md")
        elif ext in (".html", ".htm"):
            kind, dest = "design html", stem_out.with_suffix(".html")
        elif ext == ".xlsx":
            kind, dest = "xlsx → markdown (row numbers)", stem_out.with_suffix(".md")
        elif ext in IMAGE_EXT:
            kind, dest = "image", None
        elif ext in MARKITDOWN_EXT:
            kind, dest = f"{ext[1:]} → markdown", stem_out.with_suffix(".md")
        elif ext in COPY_EXT:
            kind, dest = "text", stem_out.with_suffix(ext)
        else:
            kind, dest = "unsupported", None
        rows.append((rel, src, kind, dest))
    return rows


def convert(src: Path, kind: str, dest: Path) -> str:
    """Write the converted file; return a note for the index."""
    dest.parent.mkdir(parents=True, exist_ok=True)
    if kind == "transcript":
        dest.write_text(convert_vtt(src), encoding="utf-8")
    elif kind == "design html":
        text, stubbed = convert_html(src)
        dest.write_text(text, encoding="utf-8")
        return f"{stubbed} embedded image(s) stubbed" if stubbed else ""
    elif kind == "text":
        shutil.copyfile(src, dest)
    elif src.suffix.lower() == ".xlsx":
        dest.write_text(convert_xlsx(src), encoding="utf-8")
    else:
        dest.write_text(convert_markitdown(src), encoding="utf-8")
    return ""


def main(argv=None):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--all", action="store_true", help="reconvert everything")
    ap.add_argument("--status", action="store_true", help="report only")
    args = ap.parse_args(argv)
    sys.stdout.reconfigure(encoding="utf-8")

    known = read_index()
    rows = plan()
    if not rows:
        print("inputs/ is empty. Put the BRD, transcripts, design and branding files in their subfolders first.")
        return 1

    index_rows, converted, failures = [], 0, 0
    for rel, src, kind, dest in rows:
        digest = sha256(src)
        state = "new" if rel not in known else ("changed" if known[rel][0] != digest else "unchanged")
        if args.status:
            print(f"{state:9}  {rel}")
            continue
        if not args.all and state == "unchanged" and (dest is None or dest.exists()):
            index_rows.append(known[rel][1])
            continue
        note = ""
        if dest is not None:
            try:
                note = convert(src, kind, dest)
                converted += 1
                print(f"converted  {rel}  →  {dest.relative_to(ROOT).as_posix()}")
            except Exception as e:  # keep going; report at the end
                failures += 1
                note = f"FAILED: {e}"
                print(f"FAILED     {rel}: {e}")
        elif kind == "image":
            note = "read the original — Claude views images directly"
        else:
            note = "not converted — convert by hand or remove"
        out_cell = f"`{dest.relative_to(ROOT).as_posix()}`" if dest else "—"
        size_kb = max(1, src.stat().st_size // 1024)
        index_rows.append(f"| `{rel}` | {kind} | `{digest}` | {out_cell} | {size_kb} KB | {note} |")

    if args.status:
        return 0
    OUT.mkdir(parents=True, exist_ok=True)
    INDEX.write_text(
        "# Converted inputs\n\n"
        "Generated by `uv run tools/ingest.py`. Do not edit. Skills read the converted files listed here, and the\n"
        "originals for images and visual layout. A hash that no longer matches the source means the input\n"
        "changed after conversion: re-run the tool, then re-check anything generated from it.\n\n"
        "| Source (under inputs/) | Kind | SHA-256 (16) | Converted | Size | Note |\n|---|---|---|---|---|---|\n"
        + "\n".join(index_rows)
        + "\n",
        encoding="utf-8",
    )
    print(f"{converted} converted, {len(rows) - converted - failures} unchanged or not convertible, {failures} failed. "
          "Index: inputs/_converted/INDEX.md")
    gap = check_design(rows)
    if gap:
        print(gap)
    return 1 if failures else 0


def check_design(rows) -> str:
    """The prototype's look is measured from HTML (tools/prototype_design.py) or from Figma (tools/figma_pull.py).
    Images or a PDF alone can be inventoried but not measured, so say so: /ingest records it as a question."""
    if (ROOT / "specs" / "figma.json").exists():
        return ""
    design = [src.suffix.lower() for rel, src, _, _ in rows if rel.startswith("design/")]
    if any(ext in (".html", ".htm") for ext in design):
        return ""
    has = ", ".join(sorted(set(design))) or "nothing"
    return (f"PROTOTYPE GAP  inputs/design/ has {has} but no .html, and there is no specs/figma.json. Colours, type, radii and "
            "spacing can't be measured, only estimated from images (inputs/README.md, The design): ask the client for the "
            "Figma file or the prototype's HTML, and record the question in PROJECT.md section 4.")


if __name__ == "__main__":
    sys.exit(main())
