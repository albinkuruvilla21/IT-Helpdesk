#!/usr/bin/env python3
"""
Acceptance criteria workbook — build, read, extend and check the AC .xlsx.

The workbook is the record the Product Owner signs off and the testers execute against. Claude never
hand-edits the .xlsx: it writes JSON and lets this tool lay the workbook out, so every project's
workbook has identical sheets, columns, dropdowns and formulas. Downstream skills read it back with
`dump`, never by parsing the .xlsx themselves.

  uv run tools/ac_workbook.py build   <data.json> <out.xlsx> [--force]   new workbook from JSON
  uv run tools/ac_workbook.py dump    <in.xlsx> [--format md|json] [--out FILE]
  uv run tools/ac_workbook.py append-tests <in.xlsx> <tests.json>       add test cases, keep everything else
  uv run tools/ac_workbook.py check   <in.xlsx> [--src canvas --src solution/src]   validate; trace AC IDs into the app and flows
  uv run tools/ac_workbook.py diff    <old.xlsx|old.json> <new.xlsx> [--out FILE]   what a review changed

Format and rules: docs/standards/acceptance-criteria-standard.md. JSON shape: `dump --format json`
of samples/acceptance-criteria/Sample_AcceptanceCriteria.xlsx is the worked example.
Needs openpyxl: run it with uv run, which installs it (pyproject.toml).
"""
from __future__ import annotations

import argparse
import datetime as dt
import json
import re
import sys
from pathlib import Path

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is not installed. Run the tools through uv (uv run tools/<tool>.py), or: pip install --group tools")

FORMAT_VERSION = "1.1"

# ─── Vocabulary — keep in step with docs/standards/acceptance-criteria-standard.md ───────────────
AC_TYPES = ["UI", "Functional", "Data", "Integration", "Security", "Accessibility", "Performance", "Content"]
PRIORITIES = ["Must", "Should", "Could", "Won't"]
AC_STATUSES = ["Candidate", "Draft", "In Review", "Approved", "Changed", "On Hold", "Out of Scope"]
# Names used before format 1.1. Still read, with a warning to rename them.
LEGACY_STATUSES = {"Deferred": "On Hold", "Removed": "Out of Scope"}
Q_STATUSES = ["Open", "Answered", "Closed"]
TC_TYPES = ["Positive", "Negative", "Edge", "Boundary", "Stress", "Accessibility", "UI"]
TC_STATUSES = ["Not Run", "Pass", "Fail", "Blocked"]
INPUT_STATUSES = ["Awaiting", "Received", "Not Needed"]

AC_ID = re.compile(r"^AC-[A-Z0-9]{2,10}-\d{2,3}$")
TC_ID = re.compile(r"^TC-[A-Z0-9]{2,6}-\d{3}$")
IN_ID = re.compile(r"^IN-\d{2,3}$")
Q_ID = re.compile(r"^Q-\d{2,3}$")
STD_ID = re.compile(r"^STD-[A-Z0-9]{2,8}$")
Q_REF = re.compile(r"\bQ-\d{2,3}\b")
ALTERATION = re.compile(r"^\d{4}-\d{2}-\d{2} · [^·]+ · .+")
AC_REF = re.compile(r"\bAC-[A-Z0-9]{2,10}-\d{2,3}\b")
ASSUMPTION = "ASSUMPTION"
LIVE_AC = {"Approved", "Changed"}
NOT_BUILT = {"Candidate", "On Hold", "Out of Scope", "Deferred", "Removed"}

# (json key, header, width) — header text is what the sheet shows; reading finds columns by header.
AC_COLS = [
    ("id", "AC ID", 13),
    ("screenSection", "Screen Section", 24),
    ("userRole", "User Role", 16),
    ("feature", "Feature", 26),
    ("description", "Description", 48),
    ("checks", "Acceptance Checks", 62),
    ("component", "Power Apps Component", 22),
    ("data", "Data (Lists & Fields)", 34),
    ("type", "Type", 14),
    ("priority", "Priority", 10),
    ("source", "Source", 30),
    ("clientInput", "Client Input Needed", 26),
    ("openQuestions", "Open Questions", 30),
    ("status", "Status", 12),
    ("alterations", "Alterations (dated)", 40),
]
AC_REQUIRED = ["screenSection", "userRole", "feature", "description", "checks", "component", "type", "priority", "source", "status"]

IN_COLS = [
    ("id", "Input ID", 10),
    ("acId", "AC ID", 13),
    ("item", "Item", 34),
    ("value", "Value / Link", 48),
    ("providedBy", "Provided By", 18),
    ("status", "Status", 12),
    ("notes", "Notes", 36),
]

TC_COLS = [
    ("id", "TC ID", 13),
    ("acId", "AC ID", 13),
    ("check", "Check #", 9),
    ("feature", "Feature", 24),  # formula — looked up from the AC sheet
    ("component", "Power Apps Component", 20),  # formula — looked up from the AC sheet
    ("scenario", "Test Scenario", 36),
    ("type", "Type", 13),
    ("preconditions", "Preconditions", 28),
    ("steps", "Steps", 48),
    ("expected", "Expected Result", 36),
    ("iteration", "Iteration", 10),
    ("status", "Status", 11),
    ("comments", "Comments", 30),
    ("bugRef", "Bug Ref", 12),
]
TC_FORMULA_KEYS = {"feature", "component"}

REF_COLS = [("name", "Artefact", 30), ("location", "Location", 60), ("notes", "Notes", 50)]

# Questions the PO answers. IDs are shared with PROJECT.md §4. "Current baseline" is what gets built
# if nobody answers, so an open question never blocks the build silently.
Q_COLS = [
    ("id", "Q ID", 8),
    ("acIds", "AC IDs", 18),
    ("area", "Area", 22),
    ("question", "Question", 60),
    ("whyItMatters", "Why it matters", 36),
    ("currentBaseline", "Current baseline (built if unanswered)", 40),
    ("owner", "Owner", 16),
    ("status", "Status", 11),
    ("answer", "Answer · date", 40),
]

# Acies delivery standards: they apply to every criterion and are not scope for the PO to trade away.
# Test cases may cite them (STD-A11Y, check 2) exactly like a criterion.
STD_COLS = [
    ("id", "Standard ID", 13),
    ("standard", "Standard", 30),
    ("appliesTo", "Applies to", 18),
    ("checks", "Checks", 70),
    ("basis", "Basis", 34),
]

SHEET_README = "Read Me"
SHEET_AC = "Acceptance Criteria"
SHEET_IN = "Client Inputs"
SHEET_TC = "Test Cases"
SHEET_SUMMARY = "Summary"
SHEET_REF = "References"
SHEET_Q = "Questions"
SHEET_STD = "Delivery Standards"
ALL_SHEETS = [SHEET_README, SHEET_AC, SHEET_STD, SHEET_Q, SHEET_IN, SHEET_TC, SHEET_SUMMARY, SHEET_REF]

HEADER_FILL = PatternFill("solid", fgColor="1F4E8C")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)
WRAP_TOP = Alignment(wrap_text=True, vertical="top")
MAX_ROWS = 1000  # data-validation range


# ─── Writing ─────────────────────────────────────────────────────────────────────────────────────
def _header(ws, cols):
    for i, (_, header, width) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.fill, cell.font, cell.alignment = HEADER_FILL, HEADER_FONT, WRAP_TOP
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{MAX_ROWS}"


def _dropdown(ws, cols, key, values):
    idx = [k for k, _, _ in cols].index(key) + 1
    letter = get_column_letter(idx)
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.error, dv.errorTitle = "Pick a value from the list.", "Not an allowed value"
    ws.add_data_validation(dv)
    dv.add(f"{letter}2:{letter}{MAX_ROWS}")


def _cell_value(key, value):
    if key == "checks" and isinstance(value, list):
        return "\n".join(f"{i}. {c}" for i, c in enumerate(value, start=1))
    if isinstance(value, list):
        return "\n".join(str(v) for v in value)
    return value


def _write_rows(ws, cols, rows, start_row=2, formulas=None):
    for r, row in enumerate(rows, start=start_row):
        for c, (key, _, _) in enumerate(cols, start=1):
            if formulas and key in formulas:
                value = formulas[key](r)
            else:
                value = _cell_value(key, row.get(key))
            cell = ws.cell(row=r, column=c, value=value if value != "" else None)
            cell.alignment = WRAP_TOP


def _tc_formulas(with_standards=True):
    ac, std = f"'{SHEET_AC}'", f"'{SHEET_STD}'"
    def lookup(col_letter, std_letter):
        def formula(r):
            fallback = '"⚠ Unknown AC ID"'
            if with_standards:  # a test may cite a delivery standard instead of a criterion
                fallback = f'IFERROR(INDEX({std}!${std_letter}:${std_letter},MATCH(B{r},{std}!$A:$A,0)),"⚠ Unknown AC ID")'
            return f'=IF(B{r}="","",IFERROR(INDEX({ac}!${col_letter}:${col_letter},MATCH(B{r},{ac}!$A:$A,0)),{fallback}))'
        return formula
    return {"feature": lookup("D", "B"), "component": lookup("G", "C")}


def _write_readme(ws, project):
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 100
    ws["A1"] = f"{project.get('name', 'Project')} — Acceptance Criteria"
    ws["A1"].font = TITLE_FONT
    control = [
        ("Client", project.get("client", "")),
        ("Project code", project.get("code", "")),
        ("Document version", project.get("version", "0.1")),
        ("Prepared by", project.get("preparedBy", "")),
        ("Prepared on", project.get("preparedOn", dt.date.today().isoformat())),
        ("Approved by (PO)", project.get("approvedBy", "")),
        ("Approved on", project.get("approvedOn", "")),
        ("Format version", FORMAT_VERSION),
    ]
    row = 3
    ws.cell(row=row, column=1, value="Document control").font = BOLD
    for label, value in control:
        row += 1
        ws.cell(row=row, column=1, value=label)
        ws.cell(row=row, column=2, value=value)
    notes = project.get("notes") or []  # review decisions the PO wants carried with the workbook
    if notes:
        row += 2
        ws.cell(row=row, column=1, value="Baseline notes").font = BOLD
        for note in notes:
            row += 1
            ws.cell(row=row, column=1, value=note.get("label")).alignment = WRAP_TOP
            ws.cell(row=row, column=2, value=note.get("text")).alignment = WRAP_TOP
    guide = [
        ("Sheets", ""),
        (SHEET_AC, "One row per feature. 'Acceptance Checks' is a numbered list — each check is one observable, testable outcome. Test cases cite a check by its number."),
        (SHEET_STD, "Acies delivery standards (accessibility, permissions, loading/empty/error states). They apply to every criterion. They are how Acies builds, not scope to negotiate, and are not repeated in each row. Test cases cite them like criteria."),
        (SHEET_Q, "Open questions for the PO, with the IDs used in 'Open Questions' and in PROJECT.md. 'Current baseline' is what gets built if a question stays unanswered. Answer in 'Answer · date' and set Status to Answered."),
        (SHEET_IN, "Content and decisions the client must supply (links, logos, mappings). A row stays 'Awaiting' until it arrives; tests that depend on it are 'Blocked', not 'Fail'."),
        (SHEET_TC, "One row per test case. Feature and Power Apps Component are looked up from the AC ID — '⚠ Unknown AC ID' means the ID is wrong. A re-test updates Iteration and Status in place, and the earlier result moves to Comments ('It1: Fail — BUG-014. It2: Pass')."),
        (SHEET_SUMMARY, "Calculated. Do not type into it."),
        (SHEET_REF, "Where the source artefacts live. Never put credentials or passwords here — link to where they are held instead."),
        ("", ""),
        ("Rules", ""),
        ("Working copy", "This file is the one working copy. Edit it in Excel and keep its sheets, columns and drop-downs; don't regenerate it with another tool or keep a second copy. The delivery team compares versions with 'ac_workbook.py diff'."),
        ("IDs", "AC-<PROJECT>-NN: the project code, then one running number for the whole project, from 01 (AC-CON-01, AC-CON-02, …). Test cases are TC-<AREA>-NNN, AREA being a 2–6 letter code for the screen or module (NAV, HOME), and cite the AC ID. IDs are permanent: never renumber, never reuse, never delete a row. A criterion taken out of scope keeps its row with Status 'Out of Scope'."),
        ("Source", "Every criterion names where it came from: a BRD row, a transcript timestamp, a design file. 'ASSUMPTION' means nobody said it — it must have an Open Question, and the PO must confirm it before approval. Once the PO decides, the Source says 'PO decision YYYY-MM-DD (Q-nn)' instead."),
        ("Status", "Candidate (only the BRD mentions it; not in the baseline unless the PO promotes it) · Draft → In Review → Approved · Changed (altered after approval) · On Hold (agreed for later) · Out of Scope (dropped; the row stays). Only Approved and Changed rows are built. Every change after review adds a line to 'Alterations': YYYY-MM-DD · who · what."),
        ("Type", ", ".join(AC_TYPES)),
        ("Priority", ", ".join(PRIORITIES) + " (MoSCoW)."),
        ("Test types", ", ".join(TC_TYPES)),
        ("Test status", ", ".join(TC_STATUSES)),
        ("", ""),
        ("Generated by", "tools/ac_workbook.py — see docs/standards/acceptance-criteria-standard.md in the project repository."),
    ]
    row += 2
    for label, text in guide:
        a = ws.cell(row=row, column=1, value=label or None)
        b = ws.cell(row=row, column=2, value=text or None)
        if label and not text:
            a.font = BOLD
        b.alignment = WRAP_TOP
        row += 1


def _write_summary(wb, components):
    if SHEET_SUMMARY in wb.sheetnames:
        idx = wb.sheetnames.index(SHEET_SUMMARY)
        del wb[SHEET_SUMMARY]
        ws = wb.create_sheet(SHEET_SUMMARY, idx)
    else:
        ws = wb.create_sheet(SHEET_SUMMARY)
    tc, ac = f"'{SHEET_TC}'", f"'{SHEET_AC}'"
    widths = [26, 12, 11, 10, 10, 10, 10, 11]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws["A1"] = "Test execution by component (latest result of each test case)"
    ws["A1"].font = TITLE_FONT
    heads = ["Power Apps Component", "Test Cases", "Executed", "Passed", "Failed", "Blocked", "Not Run", "Pass Rate"]
    for i, h in enumerate(heads, start=1):
        c = ws.cell(row=3, column=i, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    r = 4
    for comp in components:
        ws.cell(row=r, column=1, value=comp)
        ws.cell(row=r, column=2, value=f'=COUNTIF({tc}!$E:$E,A{r})')
        ws.cell(row=r, column=3, value=f'=B{r}-G{r}')
        ws.cell(row=r, column=4, value=f'=COUNTIFS({tc}!$E:$E,A{r},{tc}!$L:$L,"Pass")')
        ws.cell(row=r, column=5, value=f'=COUNTIFS({tc}!$E:$E,A{r},{tc}!$L:$L,"Fail")')
        ws.cell(row=r, column=6, value=f'=COUNTIFS({tc}!$E:$E,A{r},{tc}!$L:$L,"Blocked")')
        ws.cell(row=r, column=7, value=f'=COUNTIFS({tc}!$E:$E,A{r},{tc}!$L:$L,"Not Run")')
        ws.cell(row=r, column=8, value=f'=IF(C{r}=0,"",D{r}/C{r})').number_format = "0%"
        r += 1
    first, last = 4, r - 1
    ws.cell(row=r, column=1, value="Total").font = BOLD
    for col in range(2, 8):
        letter = get_column_letter(col)
        ws.cell(row=r, column=col, value=f"=SUM({letter}{first}:{letter}{last})" if last >= first else 0).font = BOLD
    ws.cell(row=r, column=8, value=f'=IF(C{r}=0,"",D{r}/C{r})').number_format = "0%"

    r += 3
    ws.cell(row=r, column=1, value="Acceptance criteria by status").font = TITLE_FONT
    r += 2
    for i, h in enumerate(["Status", "Count"], start=1):
        c = ws.cell(row=r, column=i, value=h)
        c.fill, c.font = HEADER_FILL, HEADER_FONT
    for status in AC_STATUSES:
        r += 1
        ws.cell(row=r, column=1, value=status)
        ws.cell(row=r, column=2, value=f'=COUNTIF({ac}!$N:$N,A{r})')
    r += 2
    ws.cell(row=r, column=1, value="Live criteria with no test case").font = BOLD
    ws.cell(
        row=r,
        column=2,
        value=(
            f"=SUMPRODUCT(({ac}!$A$2:$A${MAX_ROWS}<>\"\")"
            f"*((({ac}!$N$2:$N${MAX_ROWS}=\"Approved\")+({ac}!$N$2:$N${MAX_ROWS}=\"Changed\"))>0)"
            f"*(COUNTIF({tc}!$B$2:$B${MAX_ROWS},{ac}!$A$2:$A${MAX_ROWS})=0))"
        ),
    )
    ws.cell(row=r + 1, column=1, value="Run `uv run tools/ac_workbook.py check` for the list.").font = Font(italic=True)


def build(data, out: Path, force: bool):
    if out.exists() and not force:
        sys.exit(f"{out} exists. Pass --force to overwrite it (this discards hand edits — prefer append-tests or editing in Excel).")
    wb = Workbook()
    readme = wb.active
    readme.title = SHEET_README
    _write_readme(readme, data.get("project", {}))

    acs = data.get("acceptanceCriteria", [])
    ws = wb.create_sheet(SHEET_AC)
    _header(ws, AC_COLS)
    _write_rows(ws, AC_COLS, acs)
    _dropdown(ws, AC_COLS, "type", AC_TYPES)
    _dropdown(ws, AC_COLS, "priority", PRIORITIES)
    _dropdown(ws, AC_COLS, "status", AC_STATUSES)

    ws = wb.create_sheet(SHEET_STD)
    _header(ws, STD_COLS)
    _write_rows(ws, STD_COLS, data.get("standards", []))

    ws = wb.create_sheet(SHEET_Q)
    _header(ws, Q_COLS)
    _write_rows(ws, Q_COLS, data.get("questions", []))
    _dropdown(ws, Q_COLS, "status", Q_STATUSES)

    ws = wb.create_sheet(SHEET_IN)
    _header(ws, IN_COLS)
    _write_rows(ws, IN_COLS, data.get("clientInputs", []))
    _dropdown(ws, IN_COLS, "status", INPUT_STATUSES)

    ws = wb.create_sheet(SHEET_TC)
    _header(ws, TC_COLS)
    _write_rows(ws, TC_COLS, data.get("testCases", []), formulas=_tc_formulas())
    _dropdown(ws, TC_COLS, "type", TC_TYPES)
    _dropdown(ws, TC_COLS, "status", TC_STATUSES)

    _write_summary(wb, _components(acs, data.get("standards", [])))

    ws = wb.create_sheet(SHEET_REF)
    _header(ws, REF_COLS)
    _write_rows(ws, REF_COLS, data.get("references", []))

    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"Wrote {out}: {len(acs)} criteria, {len(data.get('testCases', []))} test cases.")


def _components(acs, standards=()):
    seen = []
    for a in list(acs) + [{"component": s.get("appliesTo")} for s in standards]:
        comp = (a.get("component") or "").strip()
        if comp and comp not in seen:
            seen.append(comp)
    return seen


# ─── Reading ─────────────────────────────────────────────────────────────────────────────────────
def _read_sheet(wb, name, cols):
    if name not in wb.sheetnames:
        return []
    ws = wb[name]
    header_to_key = {h: k for k, h, _ in cols}
    headers = [c.value for c in ws[1]]
    keys = [header_to_key.get(h) for h in headers]
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not values or values[0] in (None, ""):
            continue
        row = {}
        for key, value in zip(keys, values):
            if key is None:
                continue
            if isinstance(value, (dt.date, dt.datetime)):
                value = value.isoformat()[:10]
            row[key] = "" if value is None else (value.strip() if isinstance(value, str) else value)
        rows.append(row)
    return rows


def _parse_checks(text):
    if isinstance(text, list):
        return text
    checks, current = [], None
    for line in str(text or "").splitlines():
        m = re.match(r"^\s*(\d+)[.)]\s+(.*)$", line)
        if m:
            if current is not None:
                checks.append(current.strip())
            current = m.group(2)
        elif current is not None:
            current += "\n" + line
        elif line.strip():
            current = line
    if current is not None:
        checks.append(current.strip())
    return checks


def read(path: Path):
    wb = load_workbook(path)  # formulas as written; looked-up columns are recomputed below
    project = {}
    if SHEET_README in wb.sheetnames:
        labels = {
            "Client": "client", "Project code": "code", "Document version": "version", "Prepared by": "preparedBy",
            "Prepared on": "preparedOn", "Approved by (PO)": "approvedBy", "Approved on": "approvedOn",
        }
        ws = wb[SHEET_README]
        title = ws["A1"].value or ""
        project["name"] = title.replace(" — Acceptance Criteria", "")
        for a, b in ws.iter_rows(min_row=2, max_col=2, values_only=True):
            if a in labels:
                project[labels[a]] = b.isoformat()[:10] if isinstance(b, (dt.date, dt.datetime)) else (b or "")
    acs = _read_sheet(wb, SHEET_AC, AC_COLS)
    for a in acs:
        a["checks"] = _parse_checks(a.get("checks"))
    standards = _read_sheet(wb, SHEET_STD, STD_COLS)
    for s in standards:
        s["checks"] = _parse_checks(s.get("checks"))
    by_id = {a["id"]: {"feature": a.get("feature", ""), "component": a.get("component", "")} for a in acs}
    by_id.update({s["id"]: {"feature": s.get("standard", ""), "component": s.get("appliesTo", "")} for s in standards})
    tcs = _read_sheet(wb, SHEET_TC, TC_COLS)
    for t in tcs:
        ac = by_id.get(t.get("acId"))
        t["feature"] = ac["feature"] if ac else "⚠ Unknown AC ID"
        t["component"] = ac["component"] if ac else "⚠ Unknown AC ID"
    return {
        "project": project,
        "sheets": list(wb.sheetnames),
        "acceptanceCriteria": acs,
        "standards": standards,
        "questions": _read_sheet(wb, SHEET_Q, Q_COLS),
        "clientInputs": _read_sheet(wb, SHEET_IN, IN_COLS),
        "testCases": tcs,
        "references": _read_sheet(wb, SHEET_REF, REF_COLS),
    }


def to_markdown(data):
    p = data["project"]
    out = [f"# {p.get('name', 'Project')} — acceptance criteria", ""]
    out.append(f"Version {p.get('version', '?')} · approved by {p.get('approvedBy') or '—'} on {p.get('approvedOn') or '—'}")
    out.append("")
    for a in data["acceptanceCriteria"]:
        out.append(f"## {a['id']} — {a.get('feature', '')}  [{a.get('status', '')}]")
        out.append("")
        for label, key in [("Screen", "screenSection"), ("Role", "userRole"), ("Component", "component"),
                           ("Data", "data"), ("Type / priority", None), ("Source", "source"),
                           ("Client input", "clientInput"), ("Open questions", "openQuestions"), ("Alterations", "alterations")]:
            if key is None:
                out.append(f"- **{label}:** {a.get('type', '')} / {a.get('priority', '')}")
            elif a.get(key):
                out.append(f"- **{label}:** " + str(a[key]).replace("\n", "; "))
        out.append("")
        out.append(str(a.get("description", "")))
        out.append("")
        for i, c in enumerate(a.get("checks", []), start=1):
            out.append(f"{i}. " + c.replace("\n", " "))
        out.append("")
    if data.get("standards"):
        out += ["## Delivery standards", ""]
        for s in data["standards"]:
            out += [f"### {s.get('id', '')} — {s.get('standard', '')} (applies to {s.get('appliesTo', '')})", ""]
            out += [f"{i}. " + c.replace("\n", " ") for i, c in enumerate(s.get("checks", []), start=1)] + [""]
    if data.get("questions"):
        out += ["## Questions", "", "| Q | AC IDs | Question | Current baseline | Status | Answer |", "|---|---|---|---|---|---|"]
        for q in data["questions"]:
            cells = [q.get(k, "") for k in ("id", "acIds", "question", "currentBaseline", "status", "answer")]
            out.append("| " + " | ".join(str(c).replace("\n", " ").replace("|", "/") for c in cells) + " |")
        out.append("")
    if data["clientInputs"]:
        out += ["## Client inputs", "", "| ID | AC | Item | Value | Status |", "|---|---|---|---|---|"]
        for i in data["clientInputs"]:
            out.append(f"| {i.get('id', '')} | {i.get('acId', '')} | {i.get('item', '')} | {i.get('value', '')} | {i.get('status', '')} |")
        out.append("")
    if data["testCases"]:
        out += ["## Test cases", "", "| TC | AC.check | Component | Scenario | Type | Iteration | Status |", "|---|---|---|---|---|---|---|"]
        for t in data["testCases"]:
            out.append(
                f"| {t.get('id', '')} | {t.get('acId', '')}.{t.get('check', '')} | {t.get('component', '')} | "
                f"{t.get('scenario', '')} | {t.get('type', '')} | {t.get('iteration', '')} | {t.get('status', '')} |"
            )
        out.append("")
    return "\n".join(out)


# ─── Extending ───────────────────────────────────────────────────────────────────────────────────
def append_tests(path: Path, tests):
    wb = load_workbook(path)
    if SHEET_TC not in wb.sheetnames:
        sys.exit(f"{path} has no '{SHEET_TC}' sheet — not a workbook this tool built.")
    ws = wb[SHEET_TC]
    existing = {row[0] for row in ws.iter_rows(min_row=2, max_col=1, values_only=True) if row[0]}
    fresh = [t for t in tests if t.get("id") not in existing]
    skipped = [t.get("id") for t in tests if t.get("id") in existing]
    next_row = 2
    for row in ws.iter_rows(min_row=2, max_col=1):
        if row[0].value not in (None, ""):
            next_row = row[0].row + 1
    _write_rows(ws, TC_COLS, fresh, start_row=next_row, formulas=_tc_formulas(SHEET_STD in wb.sheetnames))
    acs = _read_sheet(wb, SHEET_AC, AC_COLS)
    _write_summary(wb, _components(acs, _read_sheet(wb, SHEET_STD, STD_COLS)))
    wb.save(path)
    print(f"Appended {len(fresh)} test cases to {path}." + (f" Skipped existing IDs: {', '.join(skipped)}" if skipped else ""))


# ─── Checking ────────────────────────────────────────────────────────────────────────────────────
def _norm(text):
    return re.sub(r"\s+", " ", str(text or "")).strip().lower()


def check(path: Path, src: list[Path] | None):
    data = read(path)
    errors, warnings = [], []
    acs, tcs, ins = data["acceptanceCriteria"], data["testCases"], data["clientInputs"]
    stds, qs = data["standards"], data["questions"]

    # A workbook rebuilt outside this tool loses sheets, drop-downs and formulas without saying so.
    for name in ALL_SHEETS:
        if name not in data["sheets"]:
            warnings.append(f"sheet '{name}' is missing — was the workbook regenerated outside ac_workbook.py? Rebuild it from JSON")
    for name in data["sheets"]:
        if name not in ALL_SHEETS:
            warnings.append(f"sheet '{name}' is not part of the format — move its content into the standard sheets (questions → '{SHEET_Q}')")

    q_by_id = {}
    for q in qs:
        qid = str(q.get("id", ""))
        if not Q_ID.match(qid):
            errors.append(f"{qid or '(question without ID)'}: Q ID must look like Q-01")
        if qid in q_by_id:
            errors.append(f"{qid}: duplicate Q ID")
        q_by_id[qid] = q
        if q.get("status") and q["status"] not in Q_STATUSES:
            errors.append(f"{qid}: status '{q['status']}' is not one of {', '.join(Q_STATUSES)}")
        if q.get("status") == "Answered" and not q.get("answer"):
            errors.append(f"{qid}: Answered but 'Answer · date' is empty")
    seen_text = {}
    for q in qs:
        key = _norm(q.get("question"))
        if key and key in seen_text:
            warnings.append(f"{q.get('id')}: same question as {seen_text[key]} — merge them")
        seen_text.setdefault(key, q.get("id"))

    seen = set()
    code = str(data.get("project", {}).get("code") or "").strip().upper()
    for a in acs:
        aid = a.get("id", "")
        where = aid or "(row without ID)"
        if not AC_ID.match(aid):
            errors.append(f"{where}: AC ID must look like AC-<PROJECT CODE>-01")
        elif code and aid.split("-")[1] != code:
            errors.append(f"{where}: AC ID must use the project code '{code}' (AC-{code}-NN)")
        if aid in seen:
            errors.append(f"{aid}: duplicate AC ID")
        seen.add(aid)
        for key in AC_REQUIRED:
            if not a.get(key):
                errors.append(f"{where}: '{dict((k, h) for k, h, _ in AC_COLS)[key]}' is empty")
        for key, allowed in (("type", AC_TYPES), ("priority", PRIORITIES)):
            if a.get(key) and a[key] not in allowed:
                errors.append(f"{where}: {key} '{a[key]}' is not one of {', '.join(allowed)}")
        status = a.get("status")
        if status in LEGACY_STATUSES:
            warnings.append(f"{where}: status '{status}' is the old name — use '{LEGACY_STATUSES[status]}'")
        elif status and status not in AC_STATUSES:
            errors.append(f"{where}: status '{status}' is not one of {', '.join(AC_STATUSES)}")
        refs = Q_REF.findall(str(a.get("openQuestions", "")))
        if ASSUMPTION in str(a.get("source", "")).upper():
            if not a.get("openQuestions"):
                errors.append(f"{where}: source is an ASSUMPTION but no Open Question asks the PO to confirm it")
            elif refs and q_by_id and all(q_by_id.get(r, {}).get("status") in ("Answered", "Closed") for r in refs):
                warnings.append(f"{where}: source still says ASSUMPTION but {', '.join(refs)} is answered — write 'PO decision <date> (Q-nn)' in Source")
        if qs:
            for r in refs:
                if r not in q_by_id:
                    warnings.append(f"{where}: Open Questions cites {r}, which is not on the '{SHEET_Q}' sheet")
        if status in LIVE_AC and a.get("openQuestions"):
            warnings.append(f"{where}: approved with an Open Question still recorded — resolve or move it to Alterations")
        if status == "Changed" and not a.get("alterations"):
            errors.append(f"{where}: status Changed but no dated line in Alterations")
        if status in ("On Hold", "Out of Scope", "Deferred", "Removed") and not a.get("alterations"):
            warnings.append(f"{where}: {status} with no Alterations line saying who decided and why")
        for line in str(a.get("alterations", "")).splitlines():
            if line.strip() and not ALTERATION.match(line.strip()):
                warnings.append(f"{where}: Alterations line '{line.strip()[:40]}…' should read 'YYYY-MM-DD · who · what'")
        ci = str(a.get("clientInput", ""))
        if ci and not re.fullmatch(r"\s*IN-\d{2,3}(\s*,\s*IN-\d{2,3})*\s*", ci):
            warnings.append(f"{where}: Client Input Needed should hold Input IDs (IN-01), with the detail on '{SHEET_IN}'")

    by_id = {a["id"]: a for a in acs if a.get("id")}
    std_by_id = {}
    for s in stds:
        sid = s.get("id", "")
        if not STD_ID.match(sid):
            errors.append(f"{sid or '(standard without ID)'}: Standard ID must look like STD-A11Y")
        if sid in std_by_id or sid in by_id:
            errors.append(f"{sid}: duplicate ID")
        std_by_id[sid] = s
        if not s.get("checks"):
            errors.append(f"{sid}: no checks")
    for q in qs:
        for ref in AC_REF.findall(str(q.get("acIds", ""))):
            if ref not in by_id:
                warnings.append(f"{q.get('id')}: cites {ref}, which is not a criterion")

    seen = set()
    for t in tcs:
        tid = t.get("id", "")
        if not TC_ID.match(tid):
            errors.append(f"{tid or '(test row without ID)'}: TC ID must look like TC-AREA-001")
        if tid in seen:
            errors.append(f"{tid}: duplicate TC ID")
        seen.add(tid)
        ac = by_id.get(t.get("acId")) or std_by_id.get(t.get("acId"))
        if not ac:
            errors.append(f"{tid}: AC ID '{t.get('acId')}' does not exist")
            continue
        if ac.get("status") in NOT_BUILT:
            warnings.append(f"{tid}: tests {ac['id']}, which is {ac['status']}")
        n = t.get("check")
        if n not in ("", None):
            try:
                if not 1 <= int(n) <= len(ac.get("checks", [])):
                    errors.append(f"{tid}: check #{n} — {ac['id']} has {len(ac.get('checks', []))} checks")
            except (TypeError, ValueError):
                errors.append(f"{tid}: Check # '{n}' is not a number")
        for key, allowed in (("type", TC_TYPES), ("status", TC_STATUSES)):
            if t.get(key) and t[key] not in allowed:
                errors.append(f"{tid}: {key} '{t[key]}' is not one of {', '.join(allowed)}")

    for i in ins:
        if not IN_ID.match(str(i.get("id", ""))):
            errors.append(f"{i.get('id') or '(input row without ID)'}: Input ID must look like IN-01")
        if i.get("acId") and i["acId"] not in by_id and i["acId"] not in std_by_id:
            errors.append(f"{i.get('id')}: AC ID '{i['acId']}' does not exist")

    live = [a for a in acs if a.get("status") in LIVE_AC and a.get("priority") != "Won't"]
    if tcs:
        tested = {t.get("acId") for t in tcs}
        for a in live + stds:
            if a["id"] not in tested:
                warnings.append(f"{a['id']}: live criterion or standard with no test case")

    if src:
        # Power Apps: AC IDs live in Power Fx comments in canvas YAML, and in flow step notes
        # ("description") in the solution's Workflows/*.json.
        refs = {}
        for base in src:
            files = list(base.rglob("*.pa.yaml")) + list(base.rglob("Workflows/*.json"))
            for f in files:
                for m in AC_REF.findall(f.read_text(encoding="utf-8-sig", errors="ignore")):
                    refs.setdefault(m, set()).add(f.relative_to(base.parent).as_posix())
        for ref, files in sorted(refs.items()):
            if ref not in by_id:
                errors.append(f"{ref}: referenced in code ({', '.join(sorted(files))}) but not in the workbook")
        for a in live:
            if a["id"] not in refs:
                warnings.append(f"{a['id']}: live criterion not referenced in {', '.join(str(s) for s in src)} — not implemented, or not annotated")

    print(f"{path}: {len(acs)} criteria ({len(live)} live), {len(stds)} standards, {len(qs)} questions, "
          f"{len(tcs)} test cases, {len(ins)} client inputs.")
    for e in errors:
        print(f"  ERROR    {e}")
    for w in warnings:
        print(f"  WARNING  {w}")
    print(f"{len(errors)} error(s), {len(warnings)} warning(s).")
    return 1 if errors else 0


# ─── Comparing ───────────────────────────────────────────────────────────────────────────────────
def _load_any(path: Path):
    if path.suffix.lower() == ".json":
        data = json.loads(path.read_text(encoding="utf-8"))
        present = {"acceptanceCriteria": SHEET_AC, "standards": SHEET_STD, "questions": SHEET_Q, "clientInputs": SHEET_IN,
                   "testCases": SHEET_TC, "references": SHEET_REF}
        data.setdefault("sheets", [SHEET_README, SHEET_SUMMARY] + [v for k, v in present.items() if k in data])
        for key in ("acceptanceCriteria", "standards", "questions", "clientInputs", "testCases", "references"):
            data.setdefault(key, [])
        return data
    return read(path)


def diff(old_path: Path, new_path: Path, out: Path | None):
    """What a review changed, row by row: the BA's change log, and the guard against lost rows."""
    old, new = _load_any(old_path), _load_any(new_path)
    lines, problems = [f"# Changes from `{old_path.name}` to `{new_path.name}`", ""], []
    o = {a["id"]: a for a in old["acceptanceCriteria"] if a.get("id")}
    n = {a["id"]: a for a in new["acceptanceCriteria"] if a.get("id")}

    gone = [k for k in o if k not in n]
    for k in gone:
        problems.append(f"{k} ({o[k].get('feature', '')}) was deleted. IDs are permanent: restore the row and set Status 'Out of Scope'")
    for name in ALL_SHEETS:
        if name in old["sheets"] and name not in new["sheets"]:
            problems.append(f"sheet '{name}' was dropped")
    unknown = {}
    for k, a in n.items():
        if a.get("status") and a["status"] not in AC_STATUSES and a["status"] not in LEGACY_STATUSES:
            unknown.setdefault(a["status"], []).append(k)
    for status, ids in unknown.items():
        problems.append(f"status '{status}' on {len(ids)} row(s) ({', '.join(ids[:6])}{'…' if len(ids) > 6 else ''}) "
                        f"is not in the format: use one of {', '.join(AC_STATUSES)}")

    added = [k for k in n if k not in o]
    status_moves, edited = [], []
    fields = [("feature", "Feature"), ("description", "Description"), ("component", "Component"), ("data", "Data"),
              ("type", "Type"), ("priority", "Priority"), ("source", "Source"), ("clientInput", "Client input"),
              ("openQuestions", "Open questions"), ("alterations", "Alterations")]
    for k in n:
        if k not in o:
            continue
        a, b = o[k], n[k]
        if a.get("status") != b.get("status"):
            status_moves.append((k, a.get("status", ""), b.get("status", "")))
        changes = []
        oc, nc = [_norm(c) for c in a.get("checks", [])], [_norm(c) for c in b.get("checks", [])]
        if oc != nc:
            dropped = [c for c in a.get("checks", []) if _norm(c) not in nc]
            new_c = [c for c in b.get("checks", []) if _norm(c) not in oc]
            changes.append(f"checks {len(oc)} → {len(nc)}")
            changes += [f"  − {c}" for c in dropped] + [f"  + {c}" for c in new_c]
        for key, label in fields:
            if _norm(a.get(key)) != _norm(b.get(key)):
                changes.append(f"{label}: {str(b.get(key) or '—')[:160]}")
        if changes:
            edited.append((k, changes))

    lines += ["## Needs fixing before the workbook can be used", ""]
    lines += [f"- {p}" for p in problems] or ["- Nothing."]
    lines += ["", f"## Added ({len(added)})", ""] + ([f"- {k} — {n[k].get('feature', '')} [{n[k].get('status', '')}]" for k in added] or ["- None."])
    lines += ["", f"## Status changes ({len(status_moves)})", "", "| AC | From | To |", "|---|---|---|"]
    lines += [f"| {k} | {x} | {y} |" for k, x, y in status_moves]
    lines += ["", f"## Edited rows ({len(edited)})", ""]
    for k, changes in edited:
        lines.append(f"### {k} — {n[k].get('feature', '')}")
        lines += [("- " + c) if not c.startswith("  ") else ("  " + c.strip()) for c in changes] + [""]

    oq = {q.get("id"): q for q in old.get("questions", [])}
    nq = {q.get("id"): q for q in new.get("questions", [])}
    q_lines = [f"- {k}: new — {nq[k].get('question', '')[:120]}" for k in nq if k not in oq]
    q_lines += [f"- {k}: {oq[k].get('status', '')} → {nq[k].get('status', '')} {('— ' + str(nq[k].get('answer'))[:120]) if nq[k].get('answer') else ''}"
                for k in nq if k in oq and oq[k].get("status") != nq[k].get("status")]
    q_lines += [f"- {k}: deleted" for k in oq if k not in nq]
    lines += ["## Questions", ""] + (q_lines or ["- No change."])
    o_in = {i.get("id") for i in old.get("clientInputs", [])}
    n_in = {i.get("id") for i in new.get("clientInputs", [])}
    if o_in - n_in:
        lines += ["", "## Client inputs deleted", ""] + [f"- {k}" for k in sorted(o_in - n_in)]

    text = "\n".join(lines) + "\n"
    if out:
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_text(text, encoding="utf-8")
        print(f"Wrote {out}")
    else:
        print(text)
    print(f"{len(added)} added, {len(status_moves)} status changes, {len(edited)} edited, {len(problems)} problem(s).")
    return 1 if problems else 0


# ─── CLI ─────────────────────────────────────────────────────────────────────────────────────────
def main(argv=None):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    b = sub.add_parser("build")
    b.add_argument("data", type=Path)
    b.add_argument("out", type=Path)
    b.add_argument("--force", action="store_true")
    d = sub.add_parser("dump")
    d.add_argument("xlsx", type=Path)
    d.add_argument("--format", choices=["md", "json"], default="md")
    d.add_argument("--out", type=Path)
    a = sub.add_parser("append-tests")
    a.add_argument("xlsx", type=Path)
    a.add_argument("tests", type=Path)
    c = sub.add_parser("check")
    c.add_argument("xlsx", type=Path)
    c.add_argument("--src", type=Path, action="append", help="canvas/ and solution/src (repeatable)")
    f = sub.add_parser("diff")
    f.add_argument("old", type=Path)
    f.add_argument("new", type=Path)
    f.add_argument("--out", type=Path)
    args = ap.parse_args(argv)

    if args.cmd == "build":
        build(json.loads(args.data.read_text(encoding="utf-8")), args.out, args.force)
    elif args.cmd == "dump":
        data = read(args.xlsx)
        text = json.dumps(data, indent=2, ensure_ascii=False) if args.format == "json" else to_markdown(data)
        if args.out:
            args.out.parent.mkdir(parents=True, exist_ok=True)
            args.out.write_text(text, encoding="utf-8")
            print(f"Wrote {args.out}")
        else:
            sys.stdout.reconfigure(encoding="utf-8")
            print(text)
    elif args.cmd == "append-tests":
        payload = json.loads(args.tests.read_text(encoding="utf-8"))
        append_tests(args.xlsx, payload["testCases"] if isinstance(payload, dict) else payload)
    elif args.cmd == "check":
        sys.stdout.reconfigure(encoding="utf-8")
        return check(args.xlsx, args.src)
    elif args.cmd == "diff":
        return diff(args.old, args.new, args.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
